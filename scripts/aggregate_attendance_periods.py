#!/usr/bin/env python3
"""
Build Fall, Spring, and full-year attendance workbooks from monthly
`attendance_YYYY_MM.xlsx` exports (same folder as input).

School year: June through May.
  • Fall: June–December (calendar year = start year, e.g. 2025)
  • Spring: January–May (calendar year = start + 1, e.g. 2026)
  • Full: June (start year) through May (start + 1)

Skips Instagram Points sheets (name match). Monthly files should already
exclude hidden/Instagram if exported from the app; this is an extra guard.

Outputs (same directory as monthlies):
  • attendance_{start}-{end}_full.xlsx
  • attendance_Fall{start}.xlsx
  • attendance_Spring{end}.xlsx

After merge, applies the same summary tables as add_attendance_summaries.py.

Usage:
  pip install -r scripts/requirements-attendance-tools.txt
  python scripts/aggregate_attendance_periods.py /path/to/monthly/xlsx/folder
  python scripts/aggregate_attendance_periods.py /path/to/folder --dry-run

Requires: openpyxl>=3.1
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

_SCRIPTS_DIR = Path(__file__).resolve().parent
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

from openpyxl import Workbook, load_workbook

from attendance_excel_lib import (
    apply_summaries_to_workbook,
    collect_unique_sheet_name,
    copy_sheet_as_values,
    dedupe_unique_rows,
    find_attendee_header_row,
    fresh_sheet_title,
    is_aggregate_output_filename,
    is_instagram_sheet,
    parse_monthly_attendance_filename,
    read_attendee_rows,
    school_year_start_int,
)

SCOPE_SIGNINS = "By sign-ins (all events combined for this period)"
SCOPE_UNIQUE = "By unique attendees (combined for this period)"


def period_includes_month(period: str, sy: int, y: int, m: int) -> bool:
    if period == "full":
        return (y == sy and m >= 6) or (y == sy + 1 and m <= 5)
    if period == "fall":
        return y == sy and 6 <= m <= 12
    if period == "spring":
        return y == sy + 1 and 1 <= m <= 5
    raise ValueError(f"unknown period: {period!r}")


def gather_monthlies(directory: Path) -> list[tuple[Path, int, int]]:
    out: list[tuple[Path, int, int]] = []
    for p in directory.iterdir():
        if not p.is_file() or p.suffix.lower() != ".xlsx":
            continue
        if p.name.startswith("~$"):
            continue
        if is_aggregate_output_filename(p.name):
            continue
        parsed = parse_monthly_attendance_filename(p)
        if parsed is None:
            continue
        y, m = parsed
        out.append((p, y, m))
    return sorted(out, key=lambda t: (t[1], t[2]))


def output_path(directory: Path, period: str, sy: int) -> Path:
    end_y = sy + 1
    if period == "full":
        return directory / f"attendance_{sy}-{end_y}_full.xlsx"
    if period == "fall":
        return directory / f"attendance_Fall{sy}.xlsx"
    if period == "spring":
        return directory / f"attendance_Spring{end_y}.xlsx"
    raise ValueError(period)


def build_merged_workbook(
    selected: list[tuple[Path, int, int]],
) -> Workbook | None:
    if not selected:
        return None

    wb = Workbook()
    wb.remove(wb.active)
    used_titles: set[str] = set()
    unique_accum: list[tuple[str, str, str]] = []

    for path, _, _ in selected:
        src = load_workbook(path, read_only=False, data_only=True)
        try:
            un = collect_unique_sheet_name(src)
            if un is None:
                print(f"  Warning: skip {path.name} (no Unique Attendees sheet)", file=sys.stderr)
                continue

            for ws in src.worksheets:
                if is_instagram_sheet(ws.title):
                    continue
                if ws.title == un:
                    hr = find_attendee_header_row(ws)
                    if hr is not None:
                        unique_accum.extend(read_attendee_rows(ws, hr))
                    continue

                dest_title = fresh_sheet_title(ws.title, used_titles)
                copy_sheet_as_values(ws, wb, dest_title)
        finally:
            src.close()

    if not wb.sheetnames and not unique_accum:
        return None

    deduped = dedupe_unique_rows(unique_accum)
    deduped.sort(key=lambda r: (r[0].lower() if r[0] else "", r[2].lower(), r[1].lower()))

    u_title = fresh_sheet_title("Unique Attendees", used_titles)
    uws = wb.create_sheet(u_title)
    uws.append(["Name", "Major", "Class Year"])
    for row in deduped:
        uws.append(list(row))

    return wb


def run(directory: Path, *, dry_run: bool) -> int:
    monthlies = gather_monthlies(directory)
    if not monthlies:
        print(f"No monthly files matching attendance_YYYY_MM.xlsx in {directory}", file=sys.stderr)
        return 1

    school_years = sorted({school_year_start_int(y, m) for _, y, m in monthlies})
    errors = 0

    for sy in school_years:
        for period in ("full", "fall", "spring"):
            selected = [(p, y, m) for p, y, m in monthlies if period_includes_month(period, sy, y, m)]
            out = output_path(directory, period, sy)
            if not selected:
                print(f"  Skip (no months): {out.name}")
                continue

            if dry_run:
                names = ", ".join(f"{y:04d}-{m:02d}" for _, y, m in selected)
                print(f"Would create {out.name} from {len(selected)} month file(s) [{names}]")
                continue

            try:
                wb = build_merged_workbook(selected)
                if wb is None:
                    print(f"  Skip (empty workbook): {out.name}", file=sys.stderr)
                    errors += 1
                    continue

                ok, err = apply_summaries_to_workbook(
                    wb,
                    force=True,
                    scope_line1=SCOPE_SIGNINS,
                    scope_line2=SCOPE_UNIQUE,
                )
                if not ok:
                    print(f"  Error summarizing {out.name}: {err}", file=sys.stderr)
                    errors += 1
                    wb.close()
                    continue

                wb.save(out)
                wb.close()
                print(f"  Wrote {out.name}")
            except Exception as e:
                print(f"  Error {out.name}: {e}", file=sys.stderr)
                errors += 1

    return 1 if errors else 0


def main() -> int:
    p = argparse.ArgumentParser(
        description="Merge monthly attendance_YYYY_MM.xlsx into Fall, Spring, and full-year workbooks."
    )
    p.add_argument("directory", type=Path, help="Folder containing monthly attendance_YYYY_MM.xlsx files")
    p.add_argument("--dry-run", action="store_true", help="List outputs only; do not write files")
    args = p.parse_args()
    d = args.directory.expanduser().resolve()
    if not d.is_dir():
        print(f"Not a directory: {d}", file=sys.stderr)
        return 1
    return run(d, dry_run=args.dry_run)


if __name__ == "__main__":
    raise SystemExit(main())
