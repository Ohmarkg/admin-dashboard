"""
Shared helpers for attendance .xlsx tools (monthly exports, summaries, aggregates).
"""

from __future__ import annotations

import re
import sys
from collections import Counter
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet

INSTAGRAM_EVENT_SUBSTR = "instagram points"


def is_unique_attendees_sheet(title: str) -> bool:
    t = title.strip().lower()
    return t == "unique attendees" or t.startswith("unique attendees_")


def is_instagram_sheet(title: str) -> bool:
    t = title.strip().lower()
    return INSTAGRAM_EVENT_SUBSTR in t


def find_attendee_header_row(ws: Worksheet, scan_limit: int = 500) -> int | None:
    max_r = min(ws.max_row or 1, scan_limit)
    for r in range(1, max_r + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        c = ws.cell(r, 3).value
        if not a or not b or not c:
            continue
        if (
            str(a).strip().lower() == "name"
            and str(b).strip().lower() == "major"
            and str(c).strip().lower() == "class year"
        ):
            return r
    return None


def read_attendee_rows(ws: Worksheet, header_row: int) -> list[tuple[str, str, str]]:
    rows: list[tuple[str, str, str]] = []
    for r in range(header_row + 1, (ws.max_row or header_row) + 1):
        name = ws.cell(r, 1).value
        major = ws.cell(r, 2).value
        cy = ws.cell(r, 3).value
        if name is None and major is None and cy is None:
            continue
        ns = str(name).strip() if name is not None else ""
        ms = str(major).strip() if major is not None else ""
        cs = str(cy).strip() if cy is not None else ""
        if not ms:
            ms = "NA"
        if not cs:
            cs = "NA"
        rows.append((ns, ms, cs))
    return rows


def tally_major_class(rows: Iterable[tuple[str, str, str]]) -> tuple[Counter[str], Counter[str]]:
    majors: Counter[str] = Counter()
    years: Counter[str] = Counter()
    for _, m, y in rows:
        majors[m] += 1
        years[y] += 1
    return majors, years


def sorted_counter_items(c: Counter[str]) -> list[tuple[str, int]]:
    return sorted(c.items(), key=lambda x: (-x[1], x[0].lower()))


def pct(part: int, whole: int) -> float:
    if not whole:
        return 0.0
    return round(100.0 * part / whole, 1)


def distribution_table_rows(title: str, counts: Counter[str], total: int) -> list[list]:
    out: list[list] = [[title], ["Category", "Count", "%"]]
    for label, n in sorted_counter_items(counts):
        out.append([label, n, pct(n, total)])
    return out


def event_sheet_summary_rows(sign_ins: int, majors: Counter[str], years: Counter[str]) -> list[list]:
    lines: list[list] = [
        ["Sign-in count", sign_ins],
        [],
    ]
    lines.extend(distribution_table_rows("Major distribution", majors, sign_ins))
    lines.append([])
    lines.extend(distribution_table_rows("Class year distribution", years, sign_ins))
    return lines


def unique_sheet_summary_rows(
    weighted_majors: Counter[str],
    weighted_years: Counter[str],
    total_sign_ins: int,
    unique_majors: Counter[str],
    unique_years: Counter[str],
    unique_people: int,
    *,
    scope_line1: str = "By sign-ins (all events this month)",
    scope_line2: str = "By unique attendees (this month)",
) -> list[list]:
    lines: list[list] = [
        [scope_line1],
        ["Total sign-ins", total_sign_ins],
        [],
    ]
    lines.extend(distribution_table_rows("Major distribution (by sign-in)", weighted_majors, total_sign_ins))
    lines.append([])
    lines.extend(
        distribution_table_rows("Class year distribution (by sign-in)", weighted_years, total_sign_ins)
    )
    lines.append([])
    lines.append([scope_line2])
    lines.append(["Unique people", unique_people])
    lines.append([])
    lines.extend(distribution_table_rows("Major distribution (unique)", unique_majors, unique_people))
    lines.append([])
    lines.extend(distribution_table_rows("Class year distribution (unique)", unique_years, unique_people))
    lines.append([])
    return lines


def write_block(ws: Worksheet, start_row: int, data: list[list]) -> None:
    for i, row in enumerate(data):
        for j, val in enumerate(row, start=1):
            ws.cell(row=start_row + i, column=j, value=val)


def insert_block_at_top(ws: Worksheet, rows: list[list]) -> None:
    if not rows:
        return
    n = len(rows)
    ws.insert_rows(1, n)
    write_block(ws, 1, rows)


def sign_ins_summary_starts_cell(value: object) -> bool:
    return value == "Sign-in count"


def unique_summary_starts_cell(value: object) -> bool:
    return isinstance(value, str) and value.strip().startswith("By sign-ins (")


def workbook_looks_summarized(wb: WorkbookType) -> bool:
    for ws in wb.worksheets:
        if is_unique_attendees_sheet(ws.title):
            if unique_summary_starts_cell(ws.cell(1, 1).value):
                return True
            continue
        if sign_ins_summary_starts_cell(ws.cell(1, 1).value):
            return True
    return False


def collect_unique_sheet_name(wb: WorkbookType) -> str | None:
    names = [ws.title for ws in wb.worksheets if is_unique_attendees_sheet(ws.title)]
    if not names:
        return None
    for t in names:
        if t.strip().lower() == "unique attendees":
            return t
    return names[0]


def apply_summaries_to_workbook(
    wb: WorkbookType,
    *,
    force: bool = False,
    scope_line1: str = "By sign-ins (all events this month)",
    scope_line2: str = "By unique attendees (this month)",
) -> tuple[bool, str | None]:
    """Mutate workbook in memory. Returns (ok, error_message)."""
    unique_name = collect_unique_sheet_name(wb)
    if unique_name is None:
        return False, "no 'Unique Attendees' sheet"

    event_sheets: list[Worksheet] = [ws for ws in wb.worksheets if ws.title != unique_name]

    sheet_rows: dict[str, list[tuple[str, str, str]]] = {}
    weighted: list[tuple[str, str, str]] = []

    for ws in event_sheets:
        hr = find_attendee_header_row(ws)
        if hr is None:
            print(
                f"  Skip sheet (no Name/Major/Class Year header): {ws.title}",
                file=sys.stderr,
            )
            continue
        data = read_attendee_rows(ws, hr)
        sheet_rows[ws.title] = data
        weighted.extend(data)

    uws = wb[unique_name]
    u_hr = find_attendee_header_row(uws)
    if u_hr is None:
        return False, "unique sheet has no header row"

    unique_data = read_attendee_rows(uws, u_hr)

    if not force and workbook_looks_summarized(wb):
        return False, "already has summaries (use force to re-apply)"

    wm, wy = tally_major_class(weighted)
    um, uy = tally_major_class(unique_data)
    total_sign_ins = len(weighted)
    unique_people = len(unique_data)

    u_summary = unique_sheet_summary_rows(
        wm, wy, total_sign_ins, um, uy, unique_people, scope_line1=scope_line1, scope_line2=scope_line2
    )
    insert_block_at_top(uws, u_summary)

    for ws in event_sheets:
        if ws.title not in sheet_rows:
            continue
        data = sheet_rows[ws.title]
        maj, yr = tally_major_class(data)
        insert_block_at_top(ws, event_sheet_summary_rows(len(data), maj, yr))

    return True, None


# --- Monthly filename: attendance_YYYY_MM.xlsx ---

MONTHLY_FILE_RE = re.compile(r"^attendance_(\d{4})_(\d{2})\.xlsx$", re.IGNORECASE)


def parse_monthly_attendance_filename(path: Path) -> tuple[int, int] | None:
    m = MONTHLY_FILE_RE.match(path.name)
    if not m:
        return None
    return int(m.group(1)), int(m.group(2))


def school_year_start_int(year: int, month: int) -> int:
    """June–May school year: return the June calendar year (e.g. May 2026 → 2025)."""
    if month >= 6:
        return year
    return year - 1


def is_aggregate_output_filename(name: str) -> bool:
    if re.match(r"^attendance_\d{4}-\d{4}_full\.xlsx$", name, re.I):
        return True
    if re.match(r"^attendance_Fall\d{4}\.xlsx$", name, re.I):
        return True
    if re.match(r"^attendance_Spring\d{4}\.xlsx$", name, re.I):
        return True
    return False


def norm_key_part(s: str) -> str:
    return " ".join(str(s).strip().split()).lower()


def unique_attendee_dedupe_key(row: tuple[str, str, str]) -> tuple[str, str, str]:
    """
    Same person if name + class year + major all match (case/whitespace normalized).
    Name first, then class year, then major per your rules.
    """
    name, major, cy = row
    return (norm_key_part(name), norm_key_part(cy), norm_key_part(major))


def dedupe_unique_rows(rows: list[tuple[str, str, str]]) -> list[tuple[str, str, str]]:
    seen: set[tuple[str, str, str]] = set()
    out: list[tuple[str, str, str]] = []
    for r in rows:
        k = unique_attendee_dedupe_key(r)
        if k in seen:
            continue
        seen.add(k)
        out.append(r)
    return out


_SHEET_BAD = re.compile(r"[\\/?*\[\]:]")


def sanitize_sheet_title(name: str, max_len: int = 31) -> str:
    cleaned = _SHEET_BAD.sub("-", name).strip()
    if not cleaned:
        cleaned = "Sheet"
    return cleaned[:max_len]


def fresh_sheet_title(requested: str, used: set[str]) -> str:
    base = sanitize_sheet_title(requested)
    if base not in used:
        used.add(base)
        return base
    n = 1
    while True:
        suffix = f"_{n}"
        head = base[: max(1, 31 - len(suffix))]
        cand = f"{head}{suffix}"
        if cand not in used:
            used.add(cand)
            return cand
        n += 1


def copy_sheet_as_values(src: Worksheet, dest_wb: WorkbookType, dest_title: str) -> None:
    """Append full grid of values (preserves row structure; sparse cells become None)."""
    ws = dest_wb.create_sheet(dest_title)
    max_r = src.max_row or 1
    max_c = max(src.max_column or 0, 3)
    for r in range(1, max_r + 1):
        ws.append([src.cell(r, c).value for c in range(1, max_c + 1)])


def preview_summaries_stats(wb: WorkbookType) -> tuple[bool, str | None, int, int, int]:
    """
    Read-only stats for dry-run. Returns
    (ok, err_message, event_sheets_with_header, total_sign_ins, unique_people).
    """
    unique_name = collect_unique_sheet_name(wb)
    if unique_name is None:
        return False, "no 'Unique Attendees' sheet", 0, 0, 0

    event_sheets: list[Worksheet] = [ws for ws in wb.worksheets if ws.title != unique_name]
    weighted: list[tuple[str, str, str]] = []
    count_headers = 0
    for ws in event_sheets:
        hr = find_attendee_header_row(ws)
        if hr is None:
            continue
        count_headers += 1
        weighted.extend(read_attendee_rows(ws, hr))

    uws = wb[unique_name]
    u_hr = find_attendee_header_row(uws)
    if u_hr is None:
        return False, "unique sheet has no header row", count_headers, len(weighted), 0

    unique_data = read_attendee_rows(uws, u_hr)
    return True, None, count_headers, len(weighted), len(unique_data)


def process_workbook_file(path: Path, *, dry_run: bool, force: bool) -> bool:
    """CLI helper: load path, apply summaries, save."""
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        if dry_run:
            ok, err, ns, si, uq = preview_summaries_stats(wb)
            if not ok:
                print(f"  Skip ({err}): {path.name}", file=sys.stderr)
                return False
            print(f"{path.name}: event_sheets={ns} total_sign_ins={si} unique={uq}")
            return True

        if not force and workbook_looks_summarized(wb):
            print(f"  Skip (already has summaries; use --force): {path.name}", file=sys.stderr)
            return False

        ok, err = apply_summaries_to_workbook(wb, force=force)
        if not ok:
            print(f"  Skip ({err}): {path.name}", file=sys.stderr)
            return False
        wb.save(path)
        print(f"  Wrote summaries: {path.name}")
        return True
    finally:
        wb.close()
